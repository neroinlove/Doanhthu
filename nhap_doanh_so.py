#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Công cụ nhập doanh số hàng tháng
File: nhap_doanh_so.py
"""

import sys
import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, date
from calendar import monthrange

# ===================== CẤU HÌNH =====================
FILE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tiendumoithang.xlsb.xlsx")
BACKUP_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "backup")

# Cấu trúc cột sheet NHÀ HÀNG (mới)
COLS_NH = ["Thời Gian", "Sáng", "Tối", "Tiền CK", "Tiền CK thuốc", "Tiền CK dụng cụ", "Tiền trả hàng", "Tổng"]

# Cấu trúc cột sheet THUỐC (mới - giữ nguyên)
COLS_THUOC = ["Ngày", "6h30-13h", "13h-17h", "17h-20h", "20-22h30", "Ck", "Tien tra them", "Tổng"]

# ===================== TIỆN ÍCH =====================

def clear():
    os.system("cls" if os.name == "nt" else "clear")

def header(title):
    clear()
    print("=" * 55)
    print(f"  🏪  NHẬP DOANH SỐ  —  {title}")
    print("=" * 55)

def backup_file():
    if not os.path.exists(BACKUP_DIR):
        os.makedirs(BACKUP_DIR)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = os.path.join(BACKUP_DIR, f"tiendumoithang_{ts}.xlsx")
    shutil.copy2(FILE_PATH, dest)
    return dest

def load_workbook():
    return openpyxl.load_workbook(FILE_PATH)

def save_workbook(wb):
    backup = backup_file()
    print(f"  💾 Đã backup → {os.path.basename(backup)}")
    wb.save(FILE_PATH)
    print("  ✅ Đã lưu file thành công!\n")

def input_so(prompt, default=None):
    """Nhập số, Enter để bỏ qua (trả về 0 hoặc default)"""
    hint = f" [{default}]" if default is not None else " [bỏ qua=0]"
    val = input(f"  {prompt}{hint}: ").strip()
    if val == "":
        return default if default is not None else 0
    try:
        return int(val.replace(",", "").replace(".", ""))
    except ValueError:
        print("  ⚠️  Nhập sai, dùng 0")
        return 0

def chon_thang():
    now = datetime.now()
    print(f"\n  Tháng hiện tại: {now.month}/{now.year}")
    thang = input("  Nhập tháng [Enter = tháng hiện tại]: ").strip()
    nam   = input(f"  Nhập năm    [Enter = {now.year}]: ").strip()
    thang = int(thang) if thang else now.month
    nam   = int(nam) if nam else now.year
    return thang, nam

def chon_ngay(thang, nam):
    today = date.today()
    max_day = monthrange(nam, thang)[1]
    default_day = today.day if (today.month == thang and today.year == nam) else 1
    print(f"\n  Tháng {thang}/{nam} có {max_day} ngày")
    ngay = input(f"  Nhập ngày [Enter = {default_day}]: ").strip()
    ngay = int(ngay) if ngay else default_day
    return max(1, min(ngay, max_day))

# =================== STYLE SHEET ====================

def style_header_cell(cell, bg_hex="2563EB"):
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=bg_hex)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_data_cell(cell):
    thin = Side(style="thin", color="EEEEEE")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cell.alignment = Alignment(horizontal="center", vertical="center")

def create_sheet_nh(wb, ten_sheet):
    """Tạo sheet mới kiểu Nhà Hàng với header chuẩn."""
    ws = wb.create_sheet(title=ten_sheet)
    ws.row_dimensions[1].height = 30

    # Header
    for col_idx, col_name in enumerate(COLS_NH, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        style_header_cell(cell)

    # Độ rộng cột
    widths = [12, 10, 10, 12, 14, 16, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "B2"
    return ws

def create_sheet_thuoc(wb, ten_sheet):
    """Tạo sheet mới kiểu Thuốc với header chuẩn."""
    ws = wb.create_sheet(title=ten_sheet)
    ws.row_dimensions[1].height = 30

    for col_idx, col_name in enumerate(COLS_THUOC, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        style_header_cell(cell, bg_hex="059669")

    widths = [12, 10, 10, 10, 12, 10, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "B2"
    return ws

# ================ TÌM SHEET ====================

def ten_sheet_nh(thang, nam):
    return f"tháng {thang}.{nam}"

def ten_sheet_thuoc(thang, nam):
    return f"thuốc tháng {thang}.{nam}"

def find_or_create_sheet(wb, ten_sheet, loai="nh"):
    if ten_sheet in wb.sheetnames:
        return wb[ten_sheet], False
    print(f"\n  ℹ️  Sheet '{ten_sheet}' chưa tồn tại → Tạo mới...")
    if loai == "nh":
        ws = create_sheet_nh(wb, ten_sheet)
    else:
        ws = create_sheet_thuoc(wb, ten_sheet)
    return ws, True

# ================ TÌM HÀNG NGÀY ====================

def find_row_for_day(ws, ngay, thang, nam):
    """Tìm hàng có ngày đã nhập, hoặc hàng trống tiếp theo."""
    label_ngay = f"{ngay}.{thang}"
    label_date = datetime(nam, thang, ngay)

    for row in ws.iter_rows(min_row=2, values_only=False):
        val = row[0].value
        if val is None:
            continue
        # So sánh dạng float (vd: 1.6), string hoặc datetime
        if isinstance(val, (datetime, date)):
            if val.day == ngay and val.month == thang:
                return row[0].row, True
        elif str(val).strip() in (label_ngay, str(float(ngay))):
            return row[0].row, True

    # Hàng trống tiếp theo
    return ws.max_row + 1, False

# =================== NHẬP NH ====================

def nhap_nha_hang():
    header("NHÀ HÀNG")
    wb = load_workbook()

    thang, nam = chon_thang()
    ten = ten_sheet_nh(thang, nam)
    ws, moi = find_or_create_sheet(wb, ten, "nh")

    ngay = chon_ngay(thang, nam)
    hang, exists = find_row_for_day(ws, ngay, thang, nam)

    if exists:
        print(f"\n  ⚠️  Ngày {ngay}/{thang}/{nam} đã có dữ liệu — sẽ ghi đè!")
        xn = input("  Tiếp tục? [y/N]: ").strip().lower()
        if xn != "y":
            print("  ↩️  Hủy.\n")
            return

    print(f"\n  📝 Nhập doanh thu ngày {ngay}/{thang}/{nam} (1.000đ, Enter=0):\n")
    sang        = input_so("Sáng (1.000đ)")
    toi         = input_so("Tối  (1.000đ)")
    ck          = input_so("Tiền CK")
    ck_thuoc    = input_so("Tiền CK thuốc")
    ck_dungcu   = input_so("Tiền CK dụng cụ")
    tra_hang    = input_so("Tiền trả hàng")
    tong        = sang + toi + ck + ck_thuoc + ck_dungcu

    print(f"""
  ┌─────────────────────────────────┐
  │  PREVIEW — {ngay}/{thang}/{nam:<18}│
  ├─────────────────────────────────┤
  │  Sáng           : {sang:>12,}  │
  │  Tối            : {toi:>12,}  │
  │  Tiền CK        : {ck:>12,}  │
  │  Tiền CK thuốc  : {ck_thuoc:>12,}  │
  │  Tiền CK dụng cụ: {ck_dungcu:>12,}  │
  │  Tiền trả hàng  : {tra_hang:>12,}  │
  ├─────────────────────────────────┤
  │  TỔNG           : {tong:>12,}  │
  └─────────────────────────────────┘""")

    xn = input("\n  Lưu vào Excel? [Y/n]: ").strip().lower()
    if xn == "n":
        print("  ↩️  Hủy.\n")
        return

    label = f"{ngay}.{thang}"
    data = [label, sang, toi, ck, ck_thuoc, ck_dungcu, tra_hang, tong]
    for col_idx, val in enumerate(data, start=1):
        cell = ws.cell(row=hang, column=col_idx, value=val)
        style_data_cell(cell)
        if col_idx == 1:
            cell.font = Font(bold=True)

    save_workbook(wb)
    input("  [Enter để tiếp tục]")

# =================== NHẬP THUỐC ====================

def nhap_thuoc():
    header("THUỐC")
    wb = load_workbook()

    thang, nam = chon_thang()
    ten = ten_sheet_thuoc(thang, nam)
    ws, moi = find_or_create_sheet(wb, ten, "thuoc")

    ngay = chon_ngay(thang, nam)
    hang, exists = find_row_for_day(ws, ngay, thang, nam)

    if exists:
        print(f"\n  ⚠️  Ngày {ngay}/{thang}/{nam} đã có dữ liệu — sẽ ghi đè!")
        xn = input("  Tiếp tục? [y/N]: ").strip().lower()
        if xn != "y":
            print("  ↩️  Hủy.\n")
            return

    print(f"\n  📝 Nhập doanh thu thuốc ngày {ngay}/{thang}/{nam} (1.000đ, Enter=0):\n")
    ca1       = input_so("6h30-13h  ")
    ca2       = input_so("13h-17h   ")
    ca3       = input_so("17h-20h   ")
    ca4       = input_so("20h-22h30 ")
    ck        = input_so("Tiền CK   ")
    tra_them  = input_so("Tiền trả thêm")
    tong      = ca1 + ca2 + ca3 + ca4 + ck

    print(f"""
  ┌─────────────────────────────────┐
  │  PREVIEW THUỐC — {ngay}/{thang}/{nam:<13}│
  ├─────────────────────────────────┤
  │  6h30-13h   : {ca1:>12,}      │
  │  13h-17h    : {ca2:>12,}      │
  │  17h-20h    : {ca3:>12,}      │
  │  20h-22h30  : {ca4:>12,}      │
  │  Tiền CK    : {ck:>12,}      │
  │  Tiền trả thêm: {tra_them:>10,}      │
  ├─────────────────────────────────┤
  │  TỔNG       : {tong:>12,}      │
  └─────────────────────────────────┘""")

    xn = input("\n  Lưu vào Excel? [Y/n]: ").strip().lower()
    if xn == "n":
        print("  ↩️  Hủy.\n")
        return

    label = f"{ngay}.{thang}"
    data = [label, ca1, ca2, ca3, ca4, ck, tra_them, tong]
    for col_idx, val in enumerate(data, start=1):
        cell = ws.cell(row=hang, column=col_idx, value=val)
        style_data_cell(cell)
        if col_idx == 1:
            cell.font = Font(bold=True)

    save_workbook(wb)
    input("  [Enter để tiếp tục]")

# =================== TẠO SHEET MỚI ====================

def tao_sheet_moi():
    header("TẠO SHEET THÁNG MỚI")
    wb = load_workbook()

    thang, nam = chon_thang()

    created = []
    for loai, ten_fn, create_fn in [
        ("Nhà hàng", ten_sheet_nh,   lambda w, n: create_sheet_nh(w, n)),
        ("Thuốc",   ten_sheet_thuoc, lambda w, n: create_sheet_thuoc(w, n)),
    ]:
        ten = ten_fn(thang, nam)
        if ten in wb.sheetnames:
            print(f"  ℹ️  Sheet '{ten}' đã tồn tại — bỏ qua.")
        else:
            create_fn(wb, ten)
            created.append(ten)
            print(f"  ✅ Đã tạo sheet '{ten}'")

    if created:
        save_workbook(wb)
    else:
        print("\n  Không có sheet nào được tạo.")

    input("  [Enter để tiếp tục]")

# =================== XEM TÓM TẮT ====================

def xem_tom_tat():
    header("XEM TÓM TẮT")
    wb = load_workbook()

    thang, nam = chon_thang()

    for loai, ten_fn, col_tong in [
        ("🏠 Nhà hàng", ten_sheet_nh,   8),
        ("💊 Thuốc",    ten_sheet_thuoc, 8),
    ]:
        ten = ten_fn(thang, nam)
        if ten not in wb.sheetnames:
            print(f"\n  {loai}: Chưa có sheet '{ten}'")
            continue

        ws = wb[ten]
        tong_thang = 0
        so_ngay = 0
        print(f"\n  {loai} — Tháng {thang}/{nam}")
        print("  " + "-" * 30)

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            tong_ngay = row[col_tong - 1]
            if isinstance(tong_ngay, (int, float)) and tong_ngay > 0:
                ngay_label = str(row[0]).split(" ")[0] if isinstance(row[0], str) else row[0]
                print(f"  Ngày {ngay_label:<8}: {int(tong_ngay):>12,} đ")
                tong_thang += tong_ngay
                so_ngay += 1

        print("  " + "-" * 30)
        print(f"  Tổng {so_ngay} ngày   : {int(tong_thang):>12,} đ")

    input("\n  [Enter để tiếp tục]")

# =================== MIGRATE SHEET CŨ ====================

def migrate_sheet_cu():
    header("CHUYỂN ĐỔI SHEET CŨ → CẤU TRÚC MỚI")
    wb = load_workbook()

    print("\n  Sheet nhà hàng cũ có cột: Thời Gian, Sáng, Trưa, Tối, Tổng,")
    print("  Tiền ck, Tiền mặt, Tiền thực tế, Tiền trên app, Dư thiếu")
    print("\n  Script sẽ:")
    print("  1. Giữ lại: Ngày, Sáng, Tối")
    print("  2. Dời 'Tiền ck' → 'Tiền CK' (cột 4)")
    print("  3. Thêm cột trống: Tiền CK thuốc, Tiền CK dụng cụ, Tiền trả hàng")
    print("  4. Tính lại Tổng = Sáng + Tối + Tiền CK")
    print("  5. XÓA các cột: Trưa, Tiền mặt, Tiền thực tế, Tiền trên app, Dư thiếu")
    print("\n  ⚠️  Dữ liệu sẽ được backup trước khi chỉnh sửa.")

    # Liệt kê các sheet NH cũ
    sheets_nh = []
    for name in wb.sheetnames:
        if name.startswith("tháng ") and "thuốc" not in name:
            sheets_nh.append(name)

    if not sheets_nh:
        print("\n  Không tìm thấy sheet nhà hàng nào.")
        input("  [Enter để tiếp tục]")
        return

    print(f"\n  Tìm thấy {len(sheets_nh)} sheet nhà hàng:")
    for i, s in enumerate(sheets_nh, 1):
        print(f"    {i}. {s}")

    xn = input("\n  Migrate tất cả? [y/N]: ").strip().lower()
    if xn != "y":
        print("  ↩️  Hủy.\n")
        input("  [Enter để tiếp tục]")
        return

    backup_file()
    print("  💾 Đã backup file gốc.")

    # Map cột cũ: Thời Gian=1, Sáng=2, Trưa=3, Tối=4, Tổng=5, Tiền ck=6, Tiền mặt=7, Thực tế=8, App=9, Dư thiếu=10
    for ten in sheets_nh:
        ws = wb[ten]
        header_row = [cell.value for cell in ws[1]]
        print(f"\n  Đang xử lý: {ten} | Header: {header_row[:10]}")

        # Đọc toàn bộ dữ liệu
        all_data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            all_data.append(row)

        # Xác định vị trí cột theo tên header
        def col_idx(name_list, search):
            for i, n in enumerate(name_list):
                if n and search.lower() in str(n).lower():
                    return i
            return None

        i_tgian = col_idx(header_row, "thời gian") or col_idx(header_row, "time")
        i_sang   = col_idx(header_row, "sáng") or col_idx(header_row, "sang")
        i_trua   = col_idx(header_row, "trưa") or col_idx(header_row, "trua")
        i_toi    = col_idx(header_row, "tối") or col_idx(header_row, "toi")
        i_ck     = col_idx(header_row, "ck") or col_idx(header_row, "tiền ck")

        # Xóa toàn bộ nội dung sheet
        ws.delete_rows(1, ws.max_row)

        # Viết header mới
        for ci, col_name in enumerate(COLS_NH, start=1):
            cell = ws.cell(row=1, column=ci, value=col_name)
            style_header_cell(cell)

        # Viết lại dữ liệu
        new_row = 2
        for row in all_data:
            tgian = row[i_tgian] if i_tgian is not None else None
            sang  = (row[i_sang] or 0) if i_sang is not None else 0
            toi   = (row[i_toi]  or 0) if i_toi  is not None else 0
            ck    = (row[i_ck]   or 0) if i_ck   is not None else 0
            tong  = sang + toi + ck

            if tgian is None and sang == 0 and toi == 0:
                continue  # bỏ hàng trống

            data = [tgian, sang, toi, ck, 0, 0, 0, tong]
            for ci, val in enumerate(data, start=1):
                cell = ws.cell(row=new_row, column=ci, value=val)
                style_data_cell(cell)
                if ci == 1:
                    cell.font = Font(bold=True)

            new_row += 1

        print(f"    ✅ Xong — {new_row - 2} dòng dữ liệu")

    wb.save(FILE_PATH)
    print(f"\n  ✅ Migrate hoàn tất! Đã lưu {FILE_PATH}")
    input("  [Enter để tiếp tục]")

# =================== MENU CHÍNH ====================

def main():
    while True:
        header("MENU CHÍNH")
        print("""
  1. 🏠  Nhập doanh thu Nhà Hàng
  2. 💊  Nhập doanh thu Thuốc
  3. 📅  Tạo sheet tháng mới
  4. 📊  Xem tóm tắt tháng
  5. 🔄  Migrate sheet cũ → cấu trúc mới
  0. 🚪  Thoát
""")
        chon = input("  Chọn [0-5]: ").strip()

        if chon == "1":
            nhap_nha_hang()
        elif chon == "2":
            nhap_thuoc()
        elif chon == "3":
            tao_sheet_moi()
        elif chon == "4":
            xem_tom_tat()
        elif chon == "5":
            migrate_sheet_cu()
        elif chon == "0":
            print("\n  👋 Tạm biệt!\n")
            sys.exit(0)
        else:
            print("  ⚠️  Chọn từ 0-5")
            import time; time.sleep(1)

if __name__ == "__main__":
    main()
