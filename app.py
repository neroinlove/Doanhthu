"""
Dashboard Doanh Thu - Flask App
================================
Web app nhập và theo dõi doanh thu quầy thuốc theo ngày/tháng.
"""

import json
import os
import re
import subprocess
import tempfile
import unicodedata
from datetime import datetime
from calendar import monthrange
from flask import Flask, render_template, request, jsonify

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False

BASE_DIR  = os.path.dirname(os.path.abspath(__file__))
DATA_DIR  = os.path.join(BASE_DIR, 'data', 'records')
OCR_BIN   = os.path.join(BASE_DIR, 'scripts', 'ocr_image')
OCR_SWIFT = os.path.join(BASE_DIR, 'scripts', 'ocr_image.swift')
os.makedirs(DATA_DIR, exist_ok=True)

APP_HOST = os.getenv("APP_HOST", "192.168.1.88")
APP_PORT = int(os.getenv("APP_PORT", "5501"))
APP_DEBUG = os.getenv("APP_DEBUG", "0").lower() in {"1", "true", "yes", "on"}
APP_METADATA_PATH = os.path.join(BASE_DIR, '.agents', 'metadata.json')

def load_app_version() -> str:
    """Đọc phiên bản hiển thị từ metadata phát hành của ứng dụng."""
    try:
        with open(APP_METADATA_PATH, 'r', encoding='utf-8') as f:
            return str(json.load(f).get('version', '—'))
    except (OSError, ValueError, json.JSONDecodeError):
        return '—'

APP_VERSION = load_app_version()

@app.context_processor
def inject_app_version():
    return {'app_version': APP_VERSION}

# ─── Helpers ────────────────────────────────────────────────────

def record_path(month: str) -> str:
    """month = 'YYYY-MM', e.g. '2025-07'"""
    return os.path.join(DATA_DIR, f"{month}.json")

def load_month(month: str) -> dict:
    path = record_path(month)
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return {
                "month": month,
                "quay_thuoc": data.get("quay_thuoc", []),
            }
    # Default empty structure
    return {
        "month": month,
        "quay_thuoc": [],   # list of daily records
    }

def save_month(month: str, data: dict):
    """Ghi JSON nguyên tử để tránh file dở dang nếu ứng dụng bị ngắt giữa lúc lưu."""
    target = record_path(month)
    fd, temp_path = tempfile.mkstemp(prefix=f".{month}-", suffix=".tmp", dir=DATA_DIR)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temp_path, target)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def safe_int(val, default=0) -> int:
    try:
        return int(str(val).replace(',', '').replace('.', '').strip())
    except (ValueError, TypeError):
        return default

def compute_tong_nh(rec: dict) -> int:
    return (safe_int(rec.get('sang'))
            + safe_int(rec.get('toi'))
            + safe_int(rec.get('tien_ck'))
            + safe_int(rec.get('tien_ck_thuoc'))
            + safe_int(rec.get('tien_ck_dungcu')))

def summary_nh(records: list) -> dict:
    if not records:
        return {"tong_thang": 0, "tb_ngay": 0, "max_ngay": 0, "so_ngay": 0}
    tongs = [r.get('tong', 0) for r in records if r.get('tong', 0) > 0]
    return {
        "tong_thang": sum(tongs),
        "tb_ngay":    round(sum(tongs) / len(tongs)) if tongs else 0,
        "max_ngay":   max(tongs) if tongs else 0,
        "so_ngay":    len(tongs)
    }

def list_available_months() -> list:
    """Trả về danh sách tháng đã có dữ liệu, sort mới nhất trước."""
    months = []
    for fname in os.listdir(DATA_DIR):
        if re.match(r'^\d{4}-\d{2}\.json$', fname):
            months.append(fname[:-5])
    return sorted(months, reverse=True)

def normalize_text(value: str) -> str:
    text = unicodedata.normalize("NFD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text).lower().strip()

def detect_import_field(texts: list) -> dict:
    joined = normalize_text(" ".join(texts))
    if "dung cu y te" in joined:
        return {"field": "tien_ck_dungcu", "label": "Tiền CK dụng cụ"}
    if "nha thuoc" in joined:
        return {"field": "tien_ck_thuoc", "label": "Tiền CK thuốc"}
    if "vi tam" in joined:
        return {"field": "tien_ck", "label": "Tiền CK (TK chính)"}
    return {"field": None, "label": None}

def parse_ocr_lines(lines: list) -> dict:
    texts = [str(item.get("text", "")) for item in lines]
    joined = "\n".join(texts)
    field_info = detect_import_field(texts)

    value_items = []
    date_items = []
    for item in lines:
        text = str(item.get("text", ""))
        item_x = float(item.get("x", 0))
        item_width = float(item.get("width", 0))
        for match in re.finditer(r'(\d+(?:[.,]\d+)?)\s*(Tr|K)\b', text, flags=re.IGNORECASE):
            raw_value = float(match.group(1).replace(',', '.'))
            unit = match.group(2).lower()
            amount = round(raw_value * (1_000_000 if unit == "tr" else 1_000))
            value_items.append({
                "value": float(match.group(1).replace(',', '.')),
                "unit": unit,
                "amount": amount,
                "text": text,
                "x": item_x + (item_width / 2),
                "y": float(item.get("y", 0)),
            })
        for match in re.finditer(r'(?<!\d)(\d{1,2})/(\d{1,2})(?!\d|/\d)', text):
            # Vision đôi khi gộp nhiều nhãn ngày vào một dòng. Chia đều vị trí
            # trong khung OCR vẫn cho mốc X đủ chính xác để neo cột biểu đồ.
            center_ratio = (match.start() + match.end()) / (2 * max(len(text), 1))
            date_items.append({
                "text": match.group(0),
                "x": item_x + (item_width * center_ratio),
            })

    axis_values = [item["value"] for item in value_items if item["unit"] == "tr" and item["x"] < 0.18]
    if not axis_values:
        tr_values = [item["value"] for item in value_items if item["unit"] == "tr"]
        if tr_values:
            axis_values = [max(tr_values)]

    total = None
    total_match = re.search(r'VND\s*([\d,\.]+)', joined, flags=re.IGNORECASE)
    if total_match:
        total = safe_int(total_match.group(1))

    end_date = None
    date_match = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', joined)
    if date_match:
        day, month, year = date_match.groups()
        end_date = f"{year}-{int(month):02d}-{int(day):02d}"

    range_days = None
    range_match = re.search(r'(\d+)\s*ngay\s*gan\s*(?:d)?ay', normalize_text(joined))
    if range_match:
        range_days = safe_int(range_match.group(1))

    value_labels = [
        {"amount": item["amount"], "x": item["x"]}
        for item in sorted(
            [item for item in value_items if item["x"] >= 0.18],
            key=lambda item: item["x"]
        )
    ]

    return {
        "texts": texts,
        "field": field_info["field"],
        "field_label": field_info["label"],
        "max_million": max(axis_values) if axis_values else None,
        "revenue_total": total,
        "end_date": end_date,
        "range_days": range_days,
        "value_labels": value_labels,
        "date_labels": sorted(date_items, key=lambda item: item["x"]),
    }


# ─── Routes ─────────────────────────────────────────────────────

@app.route('/')
def dashboard():
    return render_template('dashboard.html')

@app.route('/report')
def report():
    return render_template('report.html')


@app.route('/health')
def health():
    return jsonify({
        "status": "ok",
        "service": "doanhthu",
        "time": datetime.now().isoformat(timespec="seconds"),
        "data_dir": DATA_DIR,
    })


# ─── API: Records ────────────────────────────────────────────────

@app.route('/api/records/<month>', methods=['GET'])
def get_records(month):
    """Lấy toàn bộ dữ liệu tháng."""
    if not re.match(r'^\d{4}-\d{2}$', month):
        return jsonify({'error': 'month phải có dạng YYYY-MM'}), 400
    data = load_month(month)
    return jsonify(data)


@app.route('/api/records/<month>/quay_thuoc', methods=['POST'])
def upsert_nh(month):
    """Thêm hoặc cập nhật 1 ngày quầy thuốc."""
    if not re.match(r'^\d{4}-\d{2}$', month):
        return jsonify({'error': 'month phải có dạng YYYY-MM'}), 400

    body = request.get_json(force=True)
    ngay = str(body.get('ngay', '')).strip()
    if not ngay:
        return jsonify({'error': 'Thiếu trường ngay'}), 400

    rec = {
        "ngay":           ngay,
        "sang":           safe_int(body.get('sang')),
        "toi":            safe_int(body.get('toi')),
        "tien_ck":        safe_int(body.get('tien_ck')),
        "tien_ck_thuoc":  safe_int(body.get('tien_ck_thuoc')),
        "tien_ck_dungcu": safe_int(body.get('tien_ck_dungcu')),
        "tien_tra_hang":  safe_int(body.get('tien_tra_hang')),
    }
    rec["tong"] = compute_tong_nh(rec)

    data = load_month(month)
    existing = next((i for i, r in enumerate(data['quay_thuoc']) if r['ngay'] == ngay), None)
    if existing is not None:
        data['quay_thuoc'][existing] = rec
    else:
        data['quay_thuoc'].append(rec)

    # Sort theo ngày
    def day_key(r):
        try: return float(r['ngay'].split('.')[0])
        except: return 0
    data['quay_thuoc'].sort(key=day_key)

    save_month(month, data)
    return jsonify({'message': 'OK', 'record': rec})


@app.route('/api/records/<month>/<sheet_type>/<ngay>', methods=['DELETE'])
def delete_record(month, sheet_type, ngay):
    """Xóa 1 dòng ngày."""
    if sheet_type != 'quay_thuoc':
        return jsonify({'error': 'sheet_type phải là quay_thuoc'}), 400
    data = load_month(month)
    before = len(data[sheet_type])
    data[sheet_type] = [r for r in data[sheet_type] if r['ngay'] != ngay]
    if len(data[sheet_type]) == before:
        return jsonify({'error': 'Không tìm thấy bản ghi'}), 404
    save_month(month, data)
    return jsonify({'message': f'Đã xóa ngày {ngay}'})


# ─── API: Summary ────────────────────────────────────────────────

@app.route('/api/summary/<month>')
def get_summary(month):
    """KPI tóm tắt tháng."""
    data = load_month(month)
    return jsonify({
        "month":    month,
        "quay_thuoc": summary_nh(data['quay_thuoc']),
    })


@app.route('/api/months')
def get_months():
    """Danh sách tháng đã có dữ liệu."""
    return jsonify(list_available_months())


@app.route('/api/analyze-image', methods=['POST'])
def analyze_image():
    """Đọc text trong ảnh import để tự detect mốc biểu đồ và ngày cập nhật."""
    image_file = request.files.get('image')
    if not image_file:
        return jsonify({'error': 'Thiếu file ảnh'}), 400

    suffix = os.path.splitext(image_file.filename or '')[1] or '.png'
    fd, temp_path = tempfile.mkstemp(prefix='doanhthu-ocr-', suffix=suffix)
    os.close(fd)

    try:
        image_file.save(temp_path)

        if os.path.exists(OCR_BIN):
            cmd = [OCR_BIN, temp_path]
        else:
            cmd = ['swift', OCR_SWIFT, temp_path]

        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=20,
            check=True,
        )
        lines = json.loads(result.stdout or '[]')
        parsed = parse_ocr_lines(lines)
        parsed["lines"] = lines
        return jsonify(parsed)
    except subprocess.TimeoutExpired:
        return jsonify({'error': 'OCR quá lâu, vui lòng thử ảnh rõ hơn'}), 500
    except Exception as exc:
        return jsonify({'error': f'Không đọc được ảnh: {exc}'}), 500
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


# ─── API: Export Excel ───────────────────────────────────────────

@app.route('/api/export/<month>')
def export_excel(month):
    """Xuất file Excel tháng."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from flask import send_file
        import io

        data = load_month(month)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        thin = Side(style='thin', color='CCCCCC')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        def hdr(cell, color="2563EB"):
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        def dat(cell):
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        # Sheet Quầy Thuốc
        ws1 = wb.create_sheet("Quầy Thuốc")
        nh_headers = ["Ngày", "Sáng", "Tối", "Tiền CK", "Tiền CK thuốc",
                      "Tiền CK dụng cụ", "Tiền trả hàng", "Tổng"]
        for ci, h in enumerate(nh_headers, 1):
            hdr(ws1.cell(1, ci, h))
        for ri, rec in enumerate(data['quay_thuoc'], 2):
            vals = [rec.get('ngay'), rec.get('sang'), rec.get('toi'),
                    rec.get('tien_ck'), rec.get('tien_ck_thuoc'),
                    rec.get('tien_ck_dungcu'), rec.get('tien_tra_hang'), rec.get('tong')]
            for ci, v in enumerate(vals, 1):
                dat(ws1.cell(ri, ci, v))
        col_widths_nh = [10, 12, 12, 12, 16, 18, 16, 14]
        for i, w in enumerate(col_widths_nh, 1):
            ws1.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"doanhthu_{month}.xlsx"
        return send_file(buf, as_attachment=True,
                         download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except ImportError:
        return jsonify({'error': 'Cần cài openpyxl: pip install openpyxl'}), 500


# ─── API: Report data ────────────────────────────────────────────

@app.route('/api/report')
def report_data():
    """Dữ liệu cho trang báo cáo: 6 tháng gần nhất."""
    months = list_available_months()
    result = []
    for m in months[:6]:   # Trả về 6 tháng gần nhất
        data = load_month(m)
        sm_nh    = summary_nh(data['quay_thuoc'])
        result.append({
            "month":        m,
            "quay_thuoc":   sm_nh,
            "tong_chung":   sm_nh['tong_thang'],
            "nh_days":      data['quay_thuoc'],
        })
    return jsonify(result)


# ─── Entry point ─────────────────────────────────────────────────

if __name__ == '__main__':
    app.run(debug=APP_DEBUG, port=APP_PORT, host=APP_HOST)
